"""Each extractor, exercised on a real (tiny) file built on the fly."""

import zipfile
from email.message import EmailMessage

import pytest

from smart_ai_file_organizer import text_extractor as te


def _minimal_pdf(path, text="Invoice 4471 payment due by bank transfer"):
    """Write a valid one-page PDF with a text layer (no extra libraries needed)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    path.write_bytes(bytes(out))


def test_txt_and_csv(tmp_path):
    f = tmp_path / "a.txt"
    f.write_text("hello world", encoding="utf-8")
    c = tmp_path / "b.csv"
    c.write_text("name,amount\nacme,100", encoding="utf-8")
    assert te.extract_text(str(f)) == "hello world"
    assert "acme" in te.extract_text(str(c))


def test_txt_is_capped(tmp_path, monkeypatch):
    monkeypatch.setattr(te, "MAX_CHARS", 10)
    f = tmp_path / "big.txt"
    f.write_text("x" * 1000, encoding="utf-8")
    assert len(te.extract_text(str(f))) == 10


def test_pdf_text_layer(tmp_path):
    f = tmp_path / "doc.pdf"
    _minimal_pdf(f)
    assert "Invoice 4471" in te.extract_text(str(f))


def test_scanned_pdf_falls_back_to_ocr(tmp_path, monkeypatch):
    f = tmp_path / "scan.pdf"
    _minimal_pdf(f, text="")          # a page with no usable text layer
    monkeypatch.setattr(te, "_ocr_pdf_images", lambda reader: "OCR RECOVERED TEXT")
    assert te.extract_text(str(f)) == "OCR RECOVERED TEXT"


def test_scanned_pdf_without_ocr_returns_empty(tmp_path, monkeypatch):
    f = tmp_path / "scan.pdf"
    _minimal_pdf(f, text="")
    monkeypatch.setattr(te, "_configure_tesseract", lambda: False)
    assert te.extract_text(str(f)).strip() == ""


def test_docx(tmp_path):
    from docx import Document
    f = tmp_path / "r.docx"
    doc = Document()
    doc.add_paragraph("Professional summary of a software engineer")
    doc.save(f)
    assert "software engineer" in te.extract_text(str(f))


def test_xlsx(tmp_path):
    import openpyxl
    f = tmp_path / "b.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Quarter", "Revenue"])
    wb.active.append(["Q1", 1200])
    wb.save(f)
    text = te.extract_text(str(f))
    assert "Revenue" in text and "1200" in text


def test_pptx(tmp_path):
    from pptx import Presentation
    f = tmp_path / "s.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Neural networks overview"
    prs.save(f)
    assert "Neural networks" in te.extract_text(str(f))


def test_eml(tmp_path):
    msg = EmailMessage()
    msg["Subject"] = "Invoice attached"
    msg["From"] = "billing@example.com"
    msg.set_content("Please find the payment details below.")
    f = tmp_path / "m.eml"
    f.write_bytes(bytes(msg))
    text = te.extract_text(str(f))
    assert "Invoice attached" in text and "payment details" in text


def test_zip_reads_text_members_with_bounded_size(tmp_path):
    f = tmp_path / "a.zip"
    with zipfile.ZipFile(f, "w") as zf:
        zf.writestr("notes.txt", "contract clause " + "a" * 100_000)
        zf.writestr("image.bin", b"\x00\x01")
    text = te.extract_text(str(f))
    assert text.startswith("contract clause")
    assert len(text) <= 2000


def test_unsupported_and_corrupt_files_return_empty(tmp_path):
    odd = tmp_path / "x.unknown"
    odd.write_text("data")
    bad_pdf = tmp_path / "bad.pdf"
    bad_pdf.write_bytes(b"not a pdf")
    bad_docx = tmp_path / "bad.docx"
    bad_docx.write_bytes(b"not a docx")
    assert te.extract_text(str(odd)) == ""
    assert te.extract_text(str(bad_pdf)) == ""
    assert te.extract_text(str(bad_docx)) == ""


def test_tesseract_lookup_uses_path_first(monkeypatch):
    monkeypatch.setattr(te.shutil, "which", lambda name: "/usr/bin/tesseract")
    assert te._find_tesseract() == "/usr/bin/tesseract"


def test_tesseract_missing_is_reported_not_raised(monkeypatch):
    monkeypatch.setattr(te.shutil, "which", lambda name: None)
    monkeypatch.setattr(te.os.path, "exists", lambda p: False)
    assert te._find_tesseract() is None


def test_image_ocr_degrades_gracefully_without_tesseract(tmp_path, monkeypatch):
    from PIL import Image
    f = tmp_path / "pic.png"
    Image.new("RGB", (10, 10), "white").save(f)
    monkeypatch.setattr(te, "_configure_tesseract", lambda: False)
    assert te.extract_text(str(f)) == ""
